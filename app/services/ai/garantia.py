"""
garantia.py — Carta de reclamo de garantía (Word) para llantas dadas de baja por posible
falla de fabricación o de reencauche. Claude redacta el texto técnico con los datos de cada
llanta (causa, DOT, km recorridos vs estándar, cocada remanente); el monto lo calcula el código.
"""
import io
import json
from datetime import date

from .descargo import _MODEL, _cliente, _fecha_larga, _json_de, _s

IGV = 0.18

_SCHEMA = {
    "type": "object",
    "properties": {
        "asunto": {"type": "string"},
        "introduccion": {"type": "array", "items": {"type": "string"}},
        "fundamentos": {"type": "array", "items": {"type": "object", "properties": {
            "codigo": {"type": "string"}, "texto": {"type": "string"}},
            "required": ["codigo", "texto"], "additionalProperties": False}},
        "solicitud": {"type": "array", "items": {"type": "string"}},
        "cierre": {"type": "string"},
    },
    "required": ["asunto", "introduccion", "fundamentos", "solicitud", "cierre"],
    "additionalProperties": False,
}

_SYSTEM = """Eres Pedro Tenazoa, supervisor de Control de Neumáticos (Llantacentro) de TYMSAC, flota de transporte
pesado en Perú. Redactas cartas formales de reclamo de garantía a proveedores de neumáticos y reencauchadoras:
español formal, técnico, cordial y firme. No inventes datos: usa solo los que te dan. No afirmes que la falla es
de fabricación como un hecho probado; preséntala como "falla atribuible a ..." sustentada en los datos (km recorridos
muy por debajo del estándar, cocada remanente alta, DOT reciente, tipo de daño) y solicita la evaluación técnica."""

_PROMPT = """Datos del reclamo (JSON):
{datos}

Devuelve:
- asunto: una línea ("Reclamo de garantía – N neumáticos MARCA MEDIDA dados de baja prematuramente").
- introduccion: 1-2 párrafos (quiénes somos, qué se reclama y por qué).
- fundamentos: uno por llanta, en el mismo orden, {{codigo, texto}}: 2-3 oraciones con su daño, DOT/antigüedad,
  km recorridos vs estándar (% de rendimiento) y cocada remanente, explicando por qué es atribuible al producto
  (o al reencauche si la vida es reencauchada).
- solicitud: 2-4 puntos (evaluación técnica de los cascos, que están a disposición; reposición o nota de crédito
  proporcional a la vida no aprovechada por el monto referencial indicado; plazo de respuesta).
- cierre: una oración."""


def generar_carta(b: dict) -> tuple[bytes, dict]:
    items = b["items"]
    monto = round(sum(float(x.get("Reclamable") or x.get("PerdidaSoles") or 0) for x in items), 2)
    datos = {"destinatario": b.get("destinatario"), "notas_del_usuario": b.get("notas"),
             "monto_referencial_sin_igv": monto,
             "llantas": [{k: x.get(k) for k in ("Codigo", "Marca", "Modelo", "Medida", "Vida", "DOT", "AnioFab", "Causa",
                                               "Dano", "RTD", "KmTotal", "Estandar", "Rendimiento", "FMontaje",
                                               "FDesmontaje", "Placa", "ReclamarA", "Motivo", "Reclamable")} for x in items]}
    resp = _cliente().messages.create(
        model=_MODEL, max_tokens=8000, system=_SYSTEM,
        messages=[{"role": "user", "content": _PROMPT.format(datos=json.dumps(datos, ensure_ascii=False, indent=1))}],
        extra_body={"output_config": {"effort": "medium", "format": {"type": "json_schema", "schema": _SCHEMA}}},
    )
    t = _json_de(resp)
    return _armar(b, items, t, monto), {"monto": monto}


def _armar(b, items, t, monto) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    AZUL, AZUL_H = RGBColor(0x1F, 0x4E, 0x79), "1F4E79"
    doc = Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(2), Cm(1.8)
        s.left_margin = s.right_margin = Cm(2.2)
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(4)

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)

    def celda(cell, txt, bold=False, blanco=False, size=9):
        cell.text = ""; r = cell.paragraphs[0].add_run(_s(txt)); r.bold = bold; r.font.size = Pt(size)
        if blanco:
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    def parrafo(txt):
        p = doc.add_paragraph(_s(txt)); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    hoy = date.today()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(f"{_s(b.get('ciudad')) or 'Chiclayo'}, {_fecha_larga(hoy)}")
    doc.add_paragraph()
    p = doc.add_paragraph(); r = p.add_run("Señores:"); r.bold = True
    p = doc.add_paragraph(); r = p.add_run(_s(b.get("destinatario")) or "PROVEEDOR"); r.bold = True
    if _s(b.get("atencion")):
        doc.add_paragraph(f"Atención: {_s(b.get('atencion'))}")
    doc.add_paragraph("Presente.-")
    p = doc.add_paragraph(); r = p.add_run("Asunto: "); r.bold = True
    r = p.add_run(t["asunto"]); r.bold = True; r.font.color.rgb = AZUL
    doc.add_paragraph("De nuestra consideración:")
    for x in t["introduccion"]:
        parrafo(x)

    cab = ["Código", "Marca / Modelo", "Medida", "Vida", "DOT", "Causa de baja", "Km recorridos", "Rend. %", "Cocada rem.", "Monto ref. S/"]
    tb = doc.add_table(rows=1, cols=len(cab)); tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(cab):
        shade(tb.rows[0].cells[j], AZUL_H); celda(tb.rows[0].cells[j], h, bold=True, blanco=True, size=8)

    def n(v, dec=0):
        try:
            return f"{float(v):,.{dec}f}"
        except Exception:
            return "—"
    for x in items:
        cs = tb.add_row().cells
        vals = [x.get("Codigo"), f"{_s(x.get('Marca'))} {_s(x.get('Modelo'))}", x.get("Medida"), x.get("Vida"), x.get("DOT"),
                x.get("Causa") or x.get("Dano"), n(x.get("KmTotal")), n(x.get("Rendimiento")), (n(x.get("RTD"), 1) + " mm") if _s(x.get("RTD")) else "—",
                n(x.get("Reclamable") or x.get("PerdidaSoles"), 2)]
        for j, v in enumerate(vals):
            celda(cs[j], v, size=8)
    cs = tb.add_row().cells
    celda(cs[0], "TOTAL", bold=True); celda(cs[-1], f"{monto:,.2f}", bold=True)
    for c in cs:
        shade(c, "EEF3F9")

    doc.add_paragraph()
    p = doc.add_paragraph(); r = p.add_run("Fundamentos técnicos"); r.bold = True; r.font.color.rgb = AZUL
    for f in t["fundamentos"]:
        p = doc.add_paragraph(style="List Bullet"); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(f"Llanta {_s(f['codigo'])}: "); r.bold = True
        p.add_run(_s(f["texto"]))
    p = doc.add_paragraph(); r = p.add_run("Solicitud"); r.bold = True; r.font.color.rgb = AZUL
    for s in t["solicitud"]:
        p = doc.add_paragraph(_s(s), style="List Number"); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p = doc.add_paragraph()
    p.add_run(f"Monto referencial reclamado: S/ {monto:,.2f} sin IGV (S/ {round(monto * (1 + IGV), 2):,.2f} con IGV), "
              "calculado en proporción a la vida útil no aprovechada de cada neumático.").italic = True
    parrafo(t["cierre"])
    doc.add_paragraph("Atentamente,"); doc.add_paragraph(); doc.add_paragraph()
    de = (_s(b.get("de")) or "Supervisor de Llantacentro")
    p = doc.add_paragraph("_______________________________")
    p = doc.add_paragraph(); r = p.add_run(de.split("–")[0].strip()); r.bold = True
    doc.add_paragraph(de.split("–")[1].strip() if "–" in de else "Control y Supervisión de Llantacentro")
    doc.add_paragraph("TYMSAC")
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()
