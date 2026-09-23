"""
descargo.py — IA para bajas de llantas (formato del Informe Técnico de Descargo TYMSAC):

1) clasificar_dano(foto): causa del daño con la lista del cuadro "CONDICIÓN ENCONTRADA" del SCRAK.
2) leer_codigo(foto): lee el código grabado en el flanco (y marca/medida/DOT si se ven).
3) generar_descargo(datos, fotos): Informe Técnico de Descargo en Word, con el mismo formato
   del informe modelo (TEU-895): encabezado, Para/De/Fecha/Asunto/Referencia, antecedentes,
   datos del equipo y de los neumáticos, análisis técnico con registro fotográfico, análisis
   económico por llanta (sin y con IGV), conclusiones, recomendaciones y firma.
   El MONTO lo calcula el código (no la IA), método TYMSAC:
       monto por llanta = (costo ÷ cocada original) × altura de salida
   Claude redacta el texto mirando las fotos.

Modelo configurable con DAMAGE_AI_MODEL (por defecto claude-opus-5).
"""
import base64
import io
import json
import os
from datetime import date

_MODEL = os.getenv("DAMAGE_AI_MODEL", "claude-opus-5")
IGV = 0.18
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]

CAUSAS = [
    "PESTAÑA DESGARRADA POR RECALENTAMIENTO",
    "CASCO SOPLADO INTERNAMENTE",
    "PASADA DE LÍMITE PARA REENCAUCHE",
    "ROTURA DE COSTADO POR GOLPES O PENETRACIONES",
    "IMPACTO EN EL COSTADO",
    "CORTES EN BANDA POR GOLPES O PENETRACIONES",
    "INCRUSTACIÓN DE OBJETO EN BANDA DE RODAMIENTO",
    "REPARACIÓN FUERA DE ESPECIFICACIÓN",
    "CASCO FATIGADO",
    "VENA EN EL COSTADO",
    "DEVUELTA DE REENCAUCHE POR PROBLEMAS DE CARCASA",
    "RODADA DESINFLADA",
    "HOMBRO SOPLADO",
    "OTRA",
]

_SCHEMA_DANO = {
    "type": "object",
    "properties": {
        "es_llanta": {"type": "boolean"},
        "causa": {"type": "string", "enum": CAUSAS},
        "descripcion_dano": {"type": "string"},
        "zona": {"type": "string", "enum": ["banda de rodamiento", "hombro", "costado", "pestaña", "interior", "no visible"]},
        "reparable": {"type": "boolean"},
        "reencauchable": {"type": "boolean"},
        "confianza": {"type": "string", "enum": ["alta", "media", "baja"]},
        "notas": {"type": "string"},
    },
    "required": ["es_llanta", "causa", "descripcion_dano", "zona", "reparable", "reencauchable", "confianza", "notas"],
    "additionalProperties": False,
}

_PROMPT_DANO = """Eres un perito de neumáticos de camión de una flota de transporte pesado en Perú.
Mira la foto del neumático dado de baja y determina la CAUSA del daño eligiendo EXACTAMENTE una de la lista permitida.

Guía de causas:
- CASCO FATIGADO: carcasa deteriorada por uso normal / DOT antiguo (más de 5 años), grietas por envejecimiento.
- IMPACTO EN EL COSTADO / ROTURA DE COSTADO: golpe contra muro, piedra o bordillo que rompe el flanco.
- VENA EN EL COSTADO: abultamiento alargado en el flanco por cuerdas rotas.
- HOMBRO SOPLADO: separación/abultamiento en el hombro (zona entre banda y costado).
- CASCO SOPLADO INTERNAMENTE: separación interna de capas, ampolla general.
- RODADA DESINFLADA: marcas circulares o caucho molido en el flanco por rodar sin presión.
- PASADA DE LÍMITE PARA REENCAUCHE: banda gastada hasta mostrar lonas o alambres.
- CORTES EN BANDA / INCRUSTACIÓN DE OBJETO: cortes o un objeto clavado en la banda de rodamiento.
Si no se puede determinar con la foto, usa "OTRA" y explica en notas.
Escribe descripcion_dano en MAYÚSCULAS, estilo del informe (ej. "IMPACTO EN EL COSTADO CON DESGARRO DE LA CARCASA").
Contexto adicional de la llanta: %s"""

_SCHEMA_COD = {
    "type": "object",
    "properties": {
        "legible": {"type": "boolean"},
        "codigo": {"type": "string", "description": "código grabado/pintado de la llanta; vacío si no se lee"},
        "marca": {"type": "string"},
        "medida": {"type": "string"},
        "dot": {"type": "string", "description": "últimos 4 dígitos del DOT (semana+año) si se ven"},
        "confianza": {"type": "string", "enum": ["alta", "media", "baja"]},
        "notas": {"type": "string"},
    },
    "required": ["legible", "codigo", "marca", "medida", "dot", "confianza", "notas"],
    "additionalProperties": False,
}

_PROMPT_COD = """Esta es la foto del flanco de un neumático de camión de la flota TYMSAC (Perú).
La empresa identifica cada llanta con un CÓDIGO propio de 4 a 6 dígitos (ej. 19076, 18887, 01040),
grabado a fuego o pintado en el flanco. Léelo con cuidado dígito por dígito.
- codigo: SOLO ese número de identificación. NO pongas la medida (11R22.5, 295/80R22.5), ni el índice de carga,
  ni el DOT, ni números del molde del fabricante. Si no lo puedes leer con seguridad, legible=false y codigo="".
- marca / medida: si se leen en la foto; si no, "".
- dot: los últimos 4 dígitos del DOT si se ven (ej. 3024); si no, "".
- notas: una frase breve (ej. "el último dígito está parcialmente borrado")."""


def _media_type(b: bytes) -> str:
    if b[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png"
    if b[:4] == b"RIFF" and b[8:12] == b"WEBP":
        return "image/webp"
    return "image/jpeg"


def _reducir(foto: bytes, lado: int = 1400) -> bytes:
    """Reduce la foto (menos tokens y un Word más liviano). Si falla, devuelve la original."""
    try:
        from PIL import Image, ImageOps
        im = ImageOps.exif_transpose(Image.open(io.BytesIO(foto)))
        im = im.convert("RGB")
        im.thumbnail((lado, lado))
        out = io.BytesIO(); im.save(out, "JPEG", quality=85)
        return out.getvalue()
    except Exception:
        return foto


def _img_block(foto: bytes) -> dict:
    return {"type": "image", "source": {"type": "base64", "media_type": _media_type(foto),
                                       "data": base64.standard_b64encode(foto).decode("utf-8")}}


def _cliente():
    import anthropic
    if not os.getenv("ANTHROPIC_API_KEY"):
        raise RuntimeError("Falta ANTHROPIC_API_KEY en el servidor")
    return anthropic.Anthropic()


def _json_de(resp) -> dict:
    if resp.stop_reason == "refusal":
        raise RuntimeError("La IA no pudo analizar esta imagen")
    txt = next((b.text for b in resp.content if b.type == "text"), "")
    return json.loads(txt)


def _vision_json(foto: bytes, prompt: str, schema: dict, max_tokens=3000) -> dict:
    resp = _cliente().messages.create(
        model=_MODEL, max_tokens=max_tokens,
        messages=[{"role": "user", "content": [_img_block(_reducir(foto)), {"type": "text", "text": prompt}]}],
        # salida JSON validada contra el esquema
        extra_body={"output_config": {"effort": "medium", "format": {"type": "json_schema", "schema": schema}}},
    )
    return _json_de(resp)


def clasificar_dano(foto: bytes, contexto: str = "") -> dict:
    return _vision_json(foto, _PROMPT_DANO % (contexto or "sin datos"), _SCHEMA_DANO)


def leer_codigo(foto: bytes) -> dict:
    r = _vision_json(foto, _PROMPT_COD, _SCHEMA_COD, 1500)
    r["codigo"] = "".join(ch for ch in (r.get("codigo") or "") if ch.isalnum())
    return r


def monto_descargo(costo: float, cocada_orig: float, cocada_retiro: float) -> dict:
    """Método TYMSAC (informe de descargo): se cobra la cocada NO aprovechada."""
    por_mm = costo / cocada_orig
    monto = round(por_mm * cocada_retiro, 2)
    return {"costo_por_mm": round(por_mm, 2), "monto": monto,
            "pct_no_aprovechado": round(cocada_retiro / cocada_orig * 100, 1),
            "pct_aprovechado": round((cocada_orig - cocada_retiro) / cocada_orig * 100, 1)}


# ── Redacción con IA ───────────────────────────────────────────────────────

_PAR = {"type": "array", "items": {"type": "string"}}
_PUNTOS = {"type": "array", "items": {"type": "object", "properties": {"titulo": {"type": "string"}, "texto": {"type": "string"}},
                                      "required": ["titulo", "texto"], "additionalProperties": False}}
_SCHEMA_INF = {
    "type": "object",
    "properties": {
        "titulo": {"type": "string"},
        "asunto": {"type": "string"},
        "tipo_dano": {"type": "string"},
        "antecedentes": _PAR,
        "naturaleza_dano": _PAR,
        "hallazgos": _PUNTOS,
        "conclusiones": _PUNTOS,
        "recomendaciones": _PUNTOS,
        "leyendas": _PAR,
    },
    "required": ["titulo", "asunto", "tipo_dano", "antecedentes", "naturaleza_dano", "hallazgos",
                 "conclusiones", "recomendaciones", "leyendas"],
    "additionalProperties": False,
}

_SYSTEM_INF = """Eres Pedro Tenazoa, supervisor de Control de Neumáticos (Llantacentro) de TYMSAC, flota de transporte
pesado en Perú. Redactas Informes Técnicos de Descargo de neumáticos: español formal y técnico, tercera persona,
preciso y sobrio. Nunca inventes hechos, cifras, horas ni lugares que no estén en los datos; si algo no se sabe,
no lo afirmes. Los montos ya vienen calculados: cítalos exactamente, no los recalcules. Si hay fotos, descríbelas
con lo que realmente se ve."""

_PROMPT_INF = """Redacta el texto del informe con estos datos (JSON):
{datos}

Hay {nfotos} foto(s) adjunta(s), en este orden, con estos tipos: {tipos}.

Qué devolver:
- titulo: línea en MAYÚSCULAS como "IMPACTO EN 02 NEUMÁTICOS DEL TRACTO – UNIDAD TEU-895 – SATIPO" (tipo de daño, cantidad, unidad, lugar).
- asunto: "Informe de Descargo – ..." (una oración: daño, cantidad y medida, unidad, códigos y posiciones, conductor).
- tipo_dano: frase corta (ej. "Impacto con desgarro en el costado (cara lateral)").
- antecedentes: 2-3 párrafos: qué pasó, cómo se comunicó/inspeccionó, y "El presente informe documenta técnicamente ... y determina el monto a descontar."
- naturaleza_dano: 1-2 párrafos técnicos sobre la morfología del daño en las fotos, por qué ocurre y si el casco es reparable/reencauchable.
- hallazgos: 3-5 puntos {{titulo corto, texto}} (ej. "Daño por impacto lateral, no por desgaste", "Retiro prematuro", "Condiciones de la maniobra", "Riesgo asociado").
- conclusiones: 3-5 puntos {{titulo, texto}}; incluye "Causa raíz", si son irreparables, "Retiro prematuro acreditado", la responsabilidad SOLO si los datos la sustentan, y "Monto a descontar" con el total sin IGV y con IGV.
- recomendaciones: 3-5 puntos {{titulo, texto}} prácticos (notificar al conductor y recabar descargo firmado, aplicar el descuento previa conformidad, retiro como scrap, inspección complementaria de la unidad, capacitación).
- leyendas: una leyenda breve por foto, en el mismo orden (ej. "Unidad TEU-895 en base Chiclayo.", "Neumático desmontado: desgarro extenso del costado con separación de la carcasa.")."""


def _fecha_larga(d: date) -> str:
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def generar_descargo(d: dict, fotos: list[dict] | None = None) -> tuple[bytes, dict]:
    """d: numero, para, de, ciudad, placa, tipo_unidad, conductor, lugar, momento, hecho, referencia,
    llantas=[{posicion, codigo, medida, marca, modelo, condicion, cocada_orig, cocada_retiro, costo, km}].
    fotos: [{bytes, tipo, leyenda}]. Devuelve (docx_bytes, resumen)."""
    fotos = [dict(f, bytes=_reducir(f["bytes"], 1600)) for f in (fotos or []) if f.get("bytes")]
    llantas = []
    for x in d.get("llantas") or []:
        c = monto_descargo(float(x["costo"]), float(x["cocada_orig"]), float(x["cocada_retiro"]))
        llantas.append({**x, **c})
    if not llantas:
        raise ValueError("Agrega al menos una llanta")
    # como en el informe modelo: el total se suma SIN redondear cada llanta (697.73 + 697.73 -> 1,395.45)
    total = round(sum(float(x["costo"]) / float(x["cocada_orig"]) * float(x["cocada_retiro"]) for x in llantas), 2)
    total_igv = round(total * (1 + IGV), 2)
    hoy = date.today()
    n = len(llantas)

    caso = {k: d.get(k) for k in ("placa", "tipo_unidad", "conductor", "lugar", "momento", "hecho", "referencia")}
    caso["llantas"] = [{k: x.get(k) for k in ("posicion", "codigo", "medida", "marca", "modelo", "condicion", "km",
                                              "cocada_orig", "cocada_retiro", "costo", "costo_por_mm", "monto")} for x in llantas]
    caso.update({"monto_total_sin_igv": total, "monto_total_con_igv": total_igv, "cantidad_llantas": n})
    tipos = ", ".join(f"{i + 1}) {f.get('tipo') or 'foto'}" for i, f in enumerate(fotos)) or "ninguna"
    contenido = [_img_block(f["bytes"]) for f in fotos[:8]]
    contenido.append({"type": "text", "text": _PROMPT_INF.format(datos=json.dumps(caso, ensure_ascii=False, indent=1),
                                                                 nfotos=len(fotos), tipos=tipos)})
    resp = _cliente().messages.create(
        model=_MODEL, max_tokens=12000, system=_SYSTEM_INF,
        messages=[{"role": "user", "content": contenido}],
        extra_body={"output_config": {"effort": "medium", "format": {"type": "json_schema", "schema": _SCHEMA_INF}}},
    )
    t = _json_de(resp)
    for i, f in enumerate(fotos):                     # leyenda del usuario > leyenda de la IA
        if not _s(f.get("leyenda")):
            f["leyenda"] = (t["leyendas"][i] if i < len(t["leyendas"]) else "")
    docx = _armar_word(d, llantas, fotos, t, total, total_igv, hoy)
    return docx, {"monto": total, "monto_igv": total_igv, "titulo": t["titulo"], "llantas": n}


# ── Armado del Word (formato del informe modelo) ─────────────────────────────

def _armar_word(d, llantas, fotos, t, total, total_igv, hoy) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    AZUL, AZUL_H = RGBColor(0x1F, 0x4E, 0x79), "1F4E79"
    CLARO, ORO = "EEF3F9", RGBColor(0xFF, 0xC0, 0x00)
    doc = Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(2), Cm(1.8)
        s.left_margin = s.right_margin = Cm(2.2)
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(4)

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)

    def borders(table, color="1F4E79", sz="6"):
        tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), sz)
            el.set(qn("w:color"), color); b.append(el)
        tblPr.append(b)

    def cell_txt(cell, txt, bold=False, color=None, size=10.5, align=None, italic=False):
        cell.text = ""; p = cell.paragraphs[0]; r = p.add_run(_s(txt)); r.bold = bold; r.italic = italic
        r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        if align: p.alignment = align
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p.paragraph_format.space_after = Pt(1)

    def h1(txt):
        p = doc.add_paragraph(); r = p.add_run(txt); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = AZUL
        p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)

    def h2(txt):
        p = doc.add_paragraph(); r = p.add_run(txt); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = AZUL
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(3)

    def parrafo(txt):
        p = doc.add_paragraph(_s(txt)); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def punto(titulo, texto):
        p = doc.add_paragraph(style="List Bullet"); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(_s(titulo).rstrip(".:") + ": " if _s(titulo) else ""); r.bold = True
        p.add_run(_s(texto))

    def tabla_2col(filas, ancho=(5.5, 11)):
        tb = doc.add_table(rows=0, cols=2); borders(tb); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, v in filas:
            c = tb.add_row().cells; cell_txt(c[0], k, bold=True); cell_txt(c[1], v); shade(c[0], CLARO)
            c[0].width, c[1].width = Cm(ancho[0]), Cm(ancho[1])
        return tb

    # pie de página: "TYMSAC – ... – Página X de Y"
    def campo(p, instr):
        for tipo, txt in (("begin", None), (None, instr), ("end", None)):
            r = p.add_run()
            if tipo:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tipo); r._r.append(fc)
            else:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
            r.font.size = Pt(8.5)
    fp = doc.sections[0].footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("TYMSAC – Control y Supervisión de Neumáticos / Llantacentro   –   Página "); r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    campo(fp, "PAGE"); r = fp.add_run(" de "); r.font.size = Pt(8.5); campo(fp, "NUMPAGES")

    # encabezado (franja azul)
    placa = _s(d.get("placa")); ciudad = _s(d.get("ciudad")) or "Chiclayo"
    numero = _s(d.get("numero")) or f"___{hoy.strftime('%d%m%y')}/TYMSAC"
    tb = doc.add_table(rows=1, cols=1); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tb.rows[0].cells[0]; shade(c, AZUL_H); c.text = ""
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"INFORME TÉCNICO DE DESCARGO N° {numero}"); r.bold = True; r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t["titulo"].upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    doc.add_paragraph()

    tabla_2col([("Para:", _s(d.get("para")) or "Gerencia General"),
                ("De:", _s(d.get("de")) or "Control y Supervisión de Llantacentro"),
                ("Fecha:", f"{ciudad}, {_fecha_larga(hoy)}"),
                ("Asunto:", t["asunto"]),
                ("Referencia:", _s(d.get("referencia")) or "Reporte fotográfico e inspección técnica de Llantacentro")])

    h1("1. ANTECEDENTES")
    for x in t["antecedentes"]:
        parrafo(x)

    h1("2. DATOS DEL EQUIPO Y DE LOS NEUMÁTICOS")
    h2("2.1 Datos del Equipo y de la Ocurrencia")
    tu = _s(d.get("tipo_unidad")) or "unidad"
    afect = f"{len(llantas):02d} unidad(es) " + ", ".join(sorted({_s(x.get('medida')) for x in llantas if _s(x.get('medida'))})) + \
            " – códigos " + ", ".join(_s(x.get("codigo")) for x in llantas) + \
            ", posiciones " + ", ".join(_s(x.get("posicion")) for x in llantas)
    tabla_2col([(f"Unidad ({tu})", placa), ("Conductor", d.get("conductor")), ("Lugar de la ocurrencia", d.get("lugar")),
                ("Momento de la ocurrencia", d.get("momento")), ("Descripción del hecho", d.get("hecho")),
                ("Neumáticos afectados", afect), ("Tipo de daño", t["tipo_dano"])])
    h2("2.2 Datos de los Neumáticos Afectados")
    cab = ["Posición", "Código", "Medida", "Condición", "Cocada original", "Altura de salida", "Costo unitario (sin IGV)"]
    tb = doc.add_table(rows=1, cols=len(cab)); borders(tb); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(cab):
        shade(tb.rows[0].cells[j], AZUL_H)
        cell_txt(tb.rows[0].cells[j], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    for x in llantas:
        vals = [x.get("posicion"), x.get("codigo"), x.get("medida"), x.get("condicion") or "",
                f"{float(x['cocada_orig']):.1f} mm", f"{float(x['cocada_retiro']):.1f} mm", f"S/ {float(x['costo']):,.2f}"]
        cs = tb.add_row().cells
        for j, v in enumerate(vals):
            cell_txt(cs[j], v, bold=j in (0, 1), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    cons = sorted({(float(x["cocada_orig"]), float(x["cocada_retiro"])) for x in llantas})
    if len(cons) == 1:
        o, s_ = cons[0]
        parrafo(f"{'Ambos neumáticos habían' if len(llantas) == 2 else ('Los neumáticos habían' if len(llantas) > 2 else 'El neumático había')} "
                f"consumido {o - s_:.1f} mm de cocada de los {o:.1f} mm originales y "
                f"{'conservaban' if len(llantas) > 1 else 'conservaba'} {s_:.1f} mm de altura de salida al momento del retiro, "
                f"es decir, {'tenían' if len(llantas) > 1 else 'tenía'} aún pendiente de aprovechar "
                f"{'la mayor parte de su' if s_ / o >= 0.5 else 'parte de su'} vida útil residual.")

    h1("3. ANÁLISIS TÉCNICO DE LA FALLA")
    h2("3.1 Naturaleza del Daño")
    for x in t["naturaleza_dano"]:
        parrafo(x)
    h2("3.2 Hallazgos Técnicos")
    for x in t["hallazgos"]:
        punto(x["titulo"], x["texto"])
    if fotos:
        h2("3.3 Registro Fotográfico")
        tb = doc.add_table(rows=0, cols=2); borders(tb); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i in range(0, len(fotos), 2):
            par = fotos[i:i + 2]
            cs = tb.add_row().cells
            if len(par) == 1:
                cs = [cs[0].merge(cs[1])]
            for j, f in enumerate(par):
                cel = cs[j]; cel.text = ""
                p = cel.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                try:
                    ancho, alto = (7.4, 6.2) if len(par) == 2 else (10.0, 7.5)
                    try:
                        from PIL import Image
                        w, h = Image.open(io.BytesIO(f["bytes"])).size
                        if h / w > alto / ancho:   # foto vertical: manda la altura
                            ancho = alto * w / h
                    except Exception:
                        pass
                    p.add_run().add_picture(io.BytesIO(f["bytes"]), width=Cm(ancho))
                except Exception:
                    p.add_run("(no se pudo insertar la foto)")
                cp = cel.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = cp.add_run(f"Foto {i + j + 1}. {_s(f.get('leyenda'))}"); rr.italic = True; rr.font.size = Pt(9)
                rr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    h1("4. ANÁLISIS ECONÓMICO")
    h2("4.1 Determinación del Monto a Descontar")
    parrafo("El monto a descontar se determina prorrateando el costo del neumático en función de la altura de salida "
            "(cocada no aprovechada), conforme al criterio aplicado por el área de Llantacentro:")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Monto por llanta = (Costo del neumático ÷ Cocada original) × Altura de salida"); r.bold = True
    r.font.color.rgb = AZUL
    cab = ["Concepto"] + [f"Pos. {_s(x.get('posicion'))} – cód. {_s(x.get('codigo'))}" for x in llantas] + (["Total"] if len(llantas) > 1 else [])
    tb = doc.add_table(rows=1, cols=len(cab)); borders(tb); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(cab):
        shade(tb.rows[0].cells[j], AZUL_H)
        cell_txt(tb.rows[0].cells[j], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10,
                 align=None if j == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    km = [x.get("km") for x in llantas]
    filas = [("Costo del neumático nuevo (sin IGV)", [f"S/ {float(x['costo']):,.2f}" for x in llantas],
              f"S/ {sum(float(x['costo']) for x in llantas):,.2f}"),
             ("Cocada original", [f"{float(x['cocada_orig']):.1f} mm" for x in llantas], "—"),
             ("Costo por milímetro", [f"S/ {x['costo_por_mm']:,.2f} / mm" for x in llantas], "—"),
             ("Altura de salida (cocada no aprovechada)", [f"{float(x['cocada_retiro']):.1f} mm" for x in llantas],
              f"{sum(float(x['cocada_retiro']) for x in llantas):.1f} mm")]
    if any(_s(k) for k in km):
        filas.append(("Kilómetros de recorrido llanta de baja",
                      [(f"{float(k):,.2f} km" if _s(k) else "—") for k in km], "—"))
    filas.append(("MONTO A DESCONTAR (sin IGV)", [f"S/ {x['monto']:,.2f}" for x in llantas], f"S/ {total:,.2f}"))
    for k, vs, tot in filas:
        cs = tb.add_row().cells; fin = k.startswith("MONTO")
        cell_txt(cs[0], k, bold=fin, size=10)
        for j, v in enumerate(vs):
            cell_txt(cs[j + 1], v, bold=fin, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        if len(llantas) > 1:
            cell_txt(cs[-1], tot, bold=fin, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        if fin:
            for cc in cs:
                shade(cc, CLARO)
    x0 = llantas[0]
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Detalle del cálculo por llanta: "); r.bold = True
    p.add_run(f"S/ {float(x0['costo']):,.2f} ÷ {float(x0['cocada_orig']):g} mm = S/ {x0['costo_por_mm']:,.2f} por milímetro; "
              f"S/ {x0['costo_por_mm']:,.2f} × {float(x0['cocada_retiro']):g} mm de altura de salida = ")
    r = p.add_run(f"S/ {x0['monto']:,.2f} por neumático"); r.bold = True
    p.add_run(". El monto total asciende a ")
    r = p.add_run(f"S/ {total:,.2f} sin IGV"); r.bold = True
    p.add_run(f" (S/ {total_igv:,.2f} con IGV).")
    h2("4.2 Resumen del Monto a Descontar")
    tb = doc.add_table(rows=1, cols=2); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0, c1 = tb.rows[0].cells; shade(c0, AZUL_H); shade(c1, AZUL_H)
    cell_txt(c0, f"MONTO TOTAL A DESCONTAR – {len(llantas):02d} NEUMÁTICO{'S' if len(llantas) > 1 else ''} (SIN IGV)", bold=True, color=ORO)
    cell_txt(c1, f"S/ {total:,.2f}", bold=True, color=ORO, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    c0.width, c1.width = Cm(12), Cm(4.5)

    h1("5. CONCLUSIONES")
    for x in t["conclusiones"]:
        punto(x["titulo"], x["texto"])
    h1("6. RECOMENDACIONES")
    for x in t["recomendaciones"]:
        punto(x["titulo"], x["texto"])

    doc.add_paragraph(); doc.add_paragraph("Atentamente,"); doc.add_paragraph()
    firma = (_s(d.get("de")) or "Supervisor de Llantacentro").split("–")[0].strip()
    p = doc.add_paragraph(); r = p.add_run(f"{firma} – Supervisor de Llantacentro"); r.bold = True
    doc.add_paragraph("TYMSAC – Control y Supervisión de Neumáticos")
    doc.add_paragraph(f"{ciudad}, {_fecha_larga(hoy)}")

    # Conformidad del conductor: la firma se inserta después (firma con el dedo en el celular)
    h1("7. CONFORMIDAD DEL CONDUCTOR")
    parrafo(f"Yo, {_s(d.get('conductor')) or '________________'}, declaro haber sido notificado del presente informe "
            f"y del monto a descontar de S/ {total:,.2f} sin IGV (S/ {total_igv:,.2f} con IGV).")
    p = doc.add_paragraph(FIRMA_MARCA); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("_______________________________"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(_s(d.get("conductor")) or "Conductor"); r.bold = True
    p = doc.add_paragraph(FECHA_FIRMA_MARCA); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


FIRMA_MARCA = "[[FIRMA_CONDUCTOR]]"
FECHA_FIRMA_MARCA = "[[FECHA_FIRMA]]"


def firmar_word(docx: bytes, firma_png: bytes | None, fecha: str = "") -> bytes:
    """Pone la firma del conductor (o deja el espacio en blanco) en el Word guardado."""
    from docx import Document
    from docx.shared import Cm
    doc = Document(io.BytesIO(docx))
    for p in doc.paragraphs:
        if p.text.strip() == FIRMA_MARCA:
            for r in p.runs:
                r.text = ""
            if firma_png:
                try:
                    p.runs[0].add_picture(io.BytesIO(firma_png), width=Cm(5))
                except Exception:
                    pass
            else:
                p.runs[0].add_break(); p.runs[0].add_break()
        elif p.text.strip() == FECHA_FIRMA_MARCA:
            for r in p.runs:
                r.text = ""
            p.runs[0].text = f"Firmado el {fecha}" if (firma_png and fecha) else "Fecha: ____ / ____ / ________"
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()
