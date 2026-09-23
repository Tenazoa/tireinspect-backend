"""
Gestión Llantacentro:
- Historial de descargos (Word guardado, firma del conductor, estado del descuento) y ranking.
- Análisis diario subido desde la PC (predicción, alertas, reencauche, garantías, rotaciones,
  pedido sugerido, resumen gerencial) para verlo en el celular.
- Reclamos de garantía: carta al proveedor/reencauchadora redactada con IA (Word).
"""
import base64
import io
import re
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ...api.deps import get_current_inspector
from ...core.database import get_db
from ...models.models import AnaliticaSnapshot, Descargo, Inspector

router = APIRouter(prefix="/gestion", tags=["gestion"])

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ESTADOS = ("generado", "firmado", "descontado", "anulado")


def _adjunto(nombre: str) -> dict:
    return {"Content-Disposition": f"attachment; filename*=UTF-8''{quote(nombre)}"}


def _sin_tildes(t: str) -> str:
    return "".join(ch for ch in unicodedata.normalize("NFD", t) if unicodedata.category(ch) != "Mn")


def _utc(d):
    return d.isoformat() + ("Z" if d and d.tzinfo is None else "") if d else None


# ── Descargos ────────────────────────────────────────────────────────────────

def _desc_dict(r: Descargo) -> dict:
    return {"id": r.id, "numero": r.numero, "placa": r.placa, "conductor": r.conductor, "titulo": r.titulo,
            "codigos": r.codigos, "llantas": r.llantas, "monto": r.monto, "monto_igv": r.monto_igv,
            "estado": r.estado, "firmado": bool(r.firma_b64), "firmado_at": _utc(r.firmado_at),
            "descontado_at": _utc(r.descontado_at), "notas": r.notas, "created_by": r.created_by,
            "created_at": _utc(r.created_at)}


def _get_desc(db: Session, did: str, inspector: Inspector) -> Descargo:
    r = db.get(Descargo, did)
    if not r or r.company_id != inspector.company_id:
        raise HTTPException(404, "Descargo no encontrado")
    return r


@router.get("/descargos")
def listar_descargos(db: Session = Depends(get_db), inspector: Inspector = Depends(get_current_inspector)):
    rs = (db.query(Descargo).filter(Descargo.company_id == inspector.company_id)
          .order_by(Descargo.created_at.desc()).all())
    items = [_desc_dict(r) for r in rs]
    vivos = [r for r in rs if r.estado != "anulado"]
    return {
        "items": items,
        "resumen": {
            "total": len(vivos),
            "monto": round(sum(r.monto or 0 for r in vivos), 2),
            "pendiente_firma": sum(1 for r in vivos if r.estado == "generado"),
            "por_descontar": round(sum(r.monto or 0 for r in vivos if r.estado != "descontado"), 2),
            "descontado": round(sum(r.monto or 0 for r in vivos if r.estado == "descontado"), 2),
        },
    }


@router.get("/descargos/{did}/word")
def descargo_word(did: str, db: Session = Depends(get_db), inspector: Inspector = Depends(get_current_inspector)):
    from ...services.ai.descargo import firmar_word
    r = _get_desc(db, did, inspector)
    if not r.docx_b64:
        raise HTTPException(404, "Este descargo no tiene Word guardado")
    firma = base64.b64decode(r.firma_b64) if r.firma_b64 else None
    fecha = r.firmado_at.strftime("%d/%m/%Y") if r.firmado_at else ""
    docx = firmar_word(base64.b64decode(r.docx_b64), firma, fecha)
    nombre = f"INFORME_DESCARGO_{(r.placa or 'SIN-PLACA').replace(' ', '')}_{(r.codigos or '').replace(', ', '_')[:40]}"
    nombre += "_FIRMADO.docx" if firma else ".docx"
    return Response(content=docx, media_type=DOCX, headers=_adjunto(nombre))


class FirmaIn(BaseModel):
    png: str            # data URL o base64 de la firma dibujada con el dedo


@router.post("/descargos/{did}/firma")
def firmar_descargo(did: str, body: FirmaIn, db: Session = Depends(get_db),
                    inspector: Inspector = Depends(get_current_inspector)):
    r = _get_desc(db, did, inspector)
    if r.estado == "anulado":
        raise HTTPException(400, "El descargo está anulado")
    b64 = body.png.split(",", 1)[1] if body.png.startswith("data:") else body.png
    try:
        png = base64.b64decode(b64, validate=True)
    except Exception:
        raise HTTPException(400, "Firma inválida")
    if len(png) < 300 or not png.startswith(b"\x89PNG"):
        raise HTTPException(400, "Firma vacía o inválida")
    if len(png) > 600_000:
        raise HTTPException(400, "La firma es demasiado grande")
    r.firma_b64 = base64.b64encode(png).decode()
    r.firmado_at = datetime.now(timezone.utc)
    if r.estado == "generado":
        r.estado = "firmado"
    db.commit()
    return _desc_dict(r)


class EstadoIn(BaseModel):
    estado: str
    notas: str | None = None


@router.post("/descargos/{did}/estado")
def estado_descargo(did: str, body: EstadoIn, db: Session = Depends(get_db),
                    inspector: Inspector = Depends(get_current_inspector)):
    if body.estado not in ESTADOS:
        raise HTTPException(400, "Estado inválido")
    r = _get_desc(db, did, inspector)
    r.estado = body.estado
    r.descontado_at = datetime.now(timezone.utc) if body.estado == "descontado" else None
    if body.estado == "generado":
        r.firma_b64, r.firmado_at = None, None
    if body.notas is not None:
        r.notas = body.notas
    db.commit()
    return _desc_dict(r)


@router.get("/ranking")
def ranking(db: Session = Depends(get_db), inspector: Inspector = Depends(get_current_inspector)):
    """Conductores (por descargos) y placas (por bajas por daño del SCRAK + descargos)."""
    rs = db.query(Descargo).filter(Descargo.company_id == inspector.company_id, Descargo.estado != "anulado").all()
    cond = defaultdict(lambda: {"descargos": 0, "llantas": 0, "monto": 0.0, "por_descontar": 0.0, "placas": set(), "ultimo": None})
    for r in rs:
        c = cond[(r.conductor or "(sin nombre)").strip().upper()]
        c["descargos"] += 1; c["llantas"] += r.llantas or 0; c["monto"] += r.monto or 0
        if r.estado != "descontado":
            c["por_descontar"] += r.monto or 0
        if r.placa:
            c["placas"].add(r.placa)
        f = r.created_at
        if f and (c["ultimo"] is None or f > c["ultimo"]):
            c["ultimo"] = f
    conductores = {k: {"conductor": k, **{**v, "monto": round(v["monto"], 2), "por_descontar": round(v["por_descontar"], 2),
                                          "placas": ", ".join(sorted(v["placas"])), "ultimo": _utc(v["ultimo"])},
                       "gasto_ruta": 0.0, "gasto_ruta_llantas": 0.0, "vales": 0}
                   for k, v in cond.items()}
    # gastos en ruta (SOLOMON, 12 meses): el nombre del descargo ("Manuel Calle Rivera") se une
    # con el de SOLOMON ("CALLE RIVERA MANUEL IGNACIO") si todas sus palabras están en él
    gr = _snap(db, inspector, "gastos_ruta")
    tok = lambda t: {w for w in re.sub(r"[^A-ZÑ ]", " ", _sin_tildes(t.upper())).split() if len(w) > 1}
    for g in ((gr.data or {}).get("por_conductor", []) if gr else []):
        nombre = str(g.get("Conductor") or "").strip().upper()
        tg = tok(nombre)
        k = next((c for c in conductores if tok(c) and tok(c) <= tg), None)
        if k is None:
            k = nombre
            conductores[k] = {"conductor": nombre, "descargos": 0, "llantas": 0, "monto": 0.0, "por_descontar": 0.0,
                              "placas": g.get("Placas") or "", "ultimo": None,
                              "gasto_ruta": 0.0, "gasto_ruta_llantas": 0.0, "vales": 0}
        c = conductores[k]
        c["conductor"] = nombre                     # nombre completo de SOLOMON
        c["gasto_ruta"] += g.get("Monto") or 0; c["gasto_ruta_llantas"] += g.get("Llantas") or 0
        c["vales"] += g.get("Vales") or 0
    conductores = sorted(conductores.values(),
                         key=lambda x: (-(x["monto"] + x["gasto_ruta_llantas"]), -x["descargos"]))
    for c in conductores:
        c["gasto_ruta"], c["gasto_ruta_llantas"] = round(c["gasto_ruta"], 2), round(c["gasto_ruta_llantas"], 2)
    snap = _snap(db, inspector, "ranking_placas")
    placas = [dict(p) for p in (snap.data or [])] if snap else []
    por_placa = defaultdict(lambda: [0, 0.0])
    for r in rs:
        if r.placa:
            por_placa[r.placa.replace("-", "")][0] += 1
            por_placa[r.placa.replace("-", "")][1] += r.monto or 0
    for p in placas:
        k = str(p.get("Placa") or "").replace("-", "")
        p["Descargos"], p["MontoDescargos"] = por_placa[k][0], round(por_placa[k][1], 2)
    return {"conductores": conductores, "placas": placas}


# ── Análisis diario (subido desde la PC) ─────────────────────────────────────

PARTES = ("gerencial", "pred", "alertas", "reencauche", "garantia", "rotacion", "pedido",
          "ranking_placas", "cpk", "meta", "gastos_ruta", "bajas")


def _snap(db: Session, inspector: Inspector, parte: str):
    return (db.query(AnaliticaSnapshot)
            .filter(AnaliticaSnapshot.company_id == inspector.company_id, AnaliticaSnapshot.parte == parte).first())


@router.post("/analitica")
def subir_analitica(body: dict, db: Session = Depends(get_db), inspector: Inspector = Depends(get_current_inspector)):
    """Recibe {parte: datos} desde subir_web.py y reemplaza cada parte."""
    ahora = datetime.now(timezone.utc)
    n = {}
    for parte, data in body.items():
        if parte not in PARTES:
            continue
        s = _snap(db, inspector, parte)
        if not s:
            s = AnaliticaSnapshot(parte=parte, company_id=inspector.company_id)
            db.add(s)
        s.data, s.updated_at = data, ahora
        n[parte] = len(data) if isinstance(data, list) else 1
    db.commit()
    return {"ok": True, "partes": n}


@router.get("/analitica/{parte}")
def ver_analitica(parte: str, db: Session = Depends(get_db), inspector: Inspector = Depends(get_current_inspector)):
    if parte not in PARTES:
        raise HTTPException(404, "Parte desconocida")
    s = _snap(db, inspector, parte)
    meta = _snap(db, inspector, "meta")
    return {"data": s.data if s else None, "updated_at": _utc(s.updated_at) if s else None,
            "meta": meta.data if meta else None}


@router.get("/pedido/excel")
def pedido_excel(db: Session = Depends(get_db), inspector: Inspector = Depends(get_current_inspector)):
    """Pedido sugerido de llantas nuevas en Excel (para cotizar)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    s = _snap(db, inspector, "pedido")
    rows = (s.data or {}).get("items", []) if s else []
    wb = Workbook(); ws = wb.active; ws.title = "Pedido sugerido"
    hoy = datetime.now().strftime("%d/%m/%Y")
    ws.append([f"PEDIDO SUGERIDO DE LLANTAS NUEVAS – TYMSAC / LLANTACENTRO – {hoy}"])
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E79")
    ws.merge_cells("A1:L1")
    ws.append(["Según la predicción de retiro (regla 1 mm cada 8,000 km, límite 3 mm) y la mejor marca por costo por km."])
    ws.merge_cells("A2:L2")
    cab = ["Medida", "Ya (vencidas)", "0-30 días", "31-60 días", "61-90 días", "Total a comprar",
           "Marca recomendada (mejor CPK)", "S/ por 1,000 km", "Precio ref. S/ (sin IGV)", "Importe ref. S/",
           "A reencauchar (90 días)", "Costo ref. reencauche S/"]
    ws.append(cab)
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2E75B6")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    tot = 0.0
    for r in rows:
        imp = (r.get("precio") or 0) * (r.get("total") or 0)
        tot += imp
        ws.append([r.get("medida"), r.get("ya", 0), r.get("d30", 0), r.get("d60", 0), r.get("d90", 0), r.get("total", 0),
                   r.get("marca") or "", r.get("cpk1000"), r.get("precio"), round(imp, 2) if imp else None,
                   r.get("reenc", 0), r.get("precio_reenc")])
    ws.append(["TOTAL", *[sum(r.get(k, 0) or 0 for r in rows) for k in ("ya", "d30", "d60", "d90", "total")],
               "", None, None, round(tot, 2), sum(r.get("reenc", 0) or 0 for r in rows), None])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="EEF3F9")
    for i, w in enumerate([14, 12, 10, 10, 10, 14, 30, 14, 18, 16, 16, 18], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for c in list(row[7:10]) + [row[11]]:
            c.number_format = "#,##0.00"
    # detalle por llanta
    ws2 = wb.create_sheet("Llantas a retirar")
    det = (s.data or {}).get("llantas", []) if s else []
    cols = ["Codigo", "Placa", "Posicion", "Marca", "Modelo", "Medida", "Vida", "CocadaHoy", "FechaRetiro", "Horizonte", "Sugerencia"]
    ws2.append(cols)
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E79")
    for r in det:
        ws2.append([r.get(k) for k in cols])
    ws2.auto_filter.ref = ws2.dimensions
    ws2.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf)
    return Response(content=buf.getvalue(), media_type=XLSX,
                    headers=_adjunto(f"PEDIDO_SUGERIDO_LLANTAS_{datetime.now():%Y%m%d}.xlsx"))


# ── Reclamos de garantía ────────────────────────────────────────────────────

class CartaIn(BaseModel):
    items: list[dict]               # filas de la parte "garantia"
    destinatario: str = ""          # proveedor / reencauchadora
    atencion: str = ""
    de: str = "Pedro Tenazoa – Control y Supervisión de Llantacentro"
    ciudad: str = "Chiclayo"
    notas: str = ""


@router.post("/garantias/carta")
def carta_garantia(body: CartaIn, inspector: Inspector = Depends(get_current_inspector)):
    from ...services.ai.garantia import generar_carta
    if not body.items:
        raise HTTPException(400, "Selecciona al menos una llanta")
    if len(body.items) > 30:
        raise HTTPException(400, "Máximo 30 llantas por carta")
    try:
        docx, res = generar_carta(body.model_dump())
    except Exception as e:
        raise HTTPException(502, f"No se pudo generar la carta: {str(e)[:200]}")
    dest = (body.destinatario or "PROVEEDOR").upper().replace(" ", "_")[:30]
    return Response(content=docx, media_type=DOCX,
                    headers={**_adjunto(f"RECLAMO_GARANTIA_{dest}_{datetime.now():%Y%m%d}.docx"),
                             "X-Descargo-Monto": str(res.get("monto", ""))})
